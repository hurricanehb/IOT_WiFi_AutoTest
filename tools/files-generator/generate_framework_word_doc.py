"""Generate PyAutoTest_IoT framework learning/sharing Word document."""
from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
OUTPUT = ROOT / "docs" / "PyAutoTest_IoT_框架学习分享.docx"


def _set_cn_font(run, name: str = "微软雅黑", size: Pt | None = None, bold: bool = False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size:
        run.font.size = size
    run.bold = bold


def add_title(doc: Document, text: str, level: int = 1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        _set_cn_font(run, size=Pt(16 if level == 1 else 14 if level == 2 else 12), bold=True)
    return p


def add_para(doc: Document, text: str, bold: bool = False, indent: bool = False):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.left_indent = Cm(0.75)
    run = p.add_run(text)
    _set_cn_font(run, size=Pt(11), bold=bold)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25
    return p


def add_bullets(doc: Document, items: list[str]):
    for item in items:
        p = doc.add_paragraph(item, style="List Bullet")
        for run in p.runs:
            _set_cn_font(run, size=Pt(11))
        p.paragraph_format.space_after = Pt(3)


def add_code(doc: Document, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    return p


def add_table(doc: Document, headers: list[str], rows: list[list[str]]):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        for p in hdr_cells[i].paragraphs:
            for run in p.runs:
                _set_cn_font(run, size=Pt(10), bold=True)
    for r_idx, row in enumerate(rows):
        cells = table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row):
            cells[c_idx].text = val
            for p in cells[c_idx].paragraphs:
                for run in p.runs:
                    _set_cn_font(run, size=Pt(10))
    doc.add_paragraph()


def build_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.8)

    # Cover
    for _ in range(4):
        doc.add_paragraph()
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title_p.add_run("PyAutoTest_IoT\n自动化测试框架\n学习分享文档")
    _set_cn_font(tr, size=Pt(26), bold=True)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run(f"文档版本：1.0    生成日期：{date.today().isoformat()}")
    _set_cn_font(sr, size=Pt(12))

    doc.add_page_break()

    # 1. 文档目的
    add_title(doc, "1. 文档目的", 1)
    add_bullets(doc, [
        "面向团队内部分享，系统介绍当前 pytest 自动化测试框架的设计思路与落地现状。",
        "统一大家对「目录分层、执行入口、数据与配置边界、报告产出」的理解。",
        "帮助新成员快速上手，并为后续扩展（新业务域、新场景、新设备）提供可复用的参考基线。",
    ])

    # 2. 框架定位
    add_title(doc, "2. 框架一句话定位", 1)
    add_para(
        doc,
        "PyAutoTest_IoT 是一个以 pytest 为执行核心、以 ADB 为设备控制通道、"
        "按业务域分层组织的 IoT 自动化测试框架。当前重点覆盖 Android WiFi / HyperWiFi 相关场景，"
        "并预留 miniAP 与通用能力扩展空间。",
    )

    # 3. 总体架构
    add_title(doc, "3. 总体架构", 1)
    add_para(doc, "框架采用「用例层 → 核心能力层 → 数据/配置层 → 执行引擎 → 报告产出」的分层设计：")
    add_code(
        doc,
        "测试代码 (testcases)\n"
        "    ├── 业务库 (core/libs)          — WiFi 控制、性能辅助、miniAP 等\n"
        "    ├── 数据加载 (core/testdata_loader)  — YAML 数据解析与查询\n"
        "    ├── 通用工具 (core/utils)       — ADB、配置、日志、报告路径\n"
        "    └── 测试数据 (testdata/*.yaml)  — SSID、国家码、弱网场景等\n"
        "\n"
        "执行入口 (run.py) → pytest 引擎 → reports/<日期>/<时间>/\n"
        "    ├── HTML 报告 (report.html)\n"
        "    ├── Allure 原始结果 (allure/raw)\n"
        "    └── 日志 (pytest.log / automation.log)",
    )
    add_para(doc, "设计原则：", bold=True)
    add_bullets(doc, [
        "测试逻辑与测试数据解耦，减少脚本硬编码。",
        "环境细节（设备序列号、密码、AP 地址等）由配置层统一管理，支持环境变量覆盖。",
        "每次执行独立报告目录，便于问题追溯与横向对比。",
    ])

    # 4. 目录结构
    add_title(doc, "4. 目录结构与职责", 1)

    add_title(doc, "4.1 根目录关键文件", 2)
    add_table(doc, ["文件", "职责"], [
        ["run.py", "推荐执行入口，封装场景选择与报告路径"],
        ["pytest.ini", "pytest 统一配置（收集规则、marker、日志格式等）"],
        ["requirements.txt", "Python 依赖定义"],
        ["README.md", "项目总览与快速使用说明"],
        [".env", "敏感环境变量（设备序列号、WiFi 密码、AI Key 等）"],
    ])

    add_title(doc, "4.2 配置层 (config/)", 2)
    add_bullets(doc, [
        "config/test_config.json：测试运行配置，含 AP/miniAP 地址、DUT 信息、Android 序列号回退值等。",
        "config/env_config.json：环境相关配置（预留）。",
        "设计原则：环境参数优先从环境变量读取，JSON 配置文件作为默认回退。",
    ])

    add_title(doc, "4.3 核心能力层 (core/)", 2)
    add_table(doc, ["子目录", "职责"], [
        ["core/utils/", "ADB 封装、配置加载、断言增强、日志、设备管理、报告路径、YAML 解析等"],
        ["core/libs/android/", "WiFi 开关/连接/断开/清理、连接检查、性能辅助等业务能力"],
        ["core/libs/access_points/", "AP 设备驱动（HaLow miniAP、传统 WiFi AP）"],
        ["core/libs/common/", "公共域能力预留"],
        ["core/testdata_loader/android/", "WiFi SSID、国家码表、弱网场景等 YAML 数据加载与查询"],
        ["core/ai/", "AI 客户端、弱网仿真、用例生成、结果分析等（可选能力）"],
    ])

    add_title(doc, "4.4 数据层 (testdata/)", 2)
    add_bullets(doc, [
        "testdata/android/wifi_ssid.yaml：WiFi SSID 列表、双 WiFi 配置、连接期望等。",
        "testdata/android/country_code_tables.yaml：国家码对照表。",
        "testdata/android/weaknet_scene.yaml：弱网仿真场景参数。",
        "testdata/miniap/：miniAP 域数据预留。",
    ])

    add_title(doc, "4.5 用例层 (testcases/)", 2)
    add_table(doc, ["目录", "说明"], [
        ["testcases/android/", "Android / HyperWiFi 主业务域（wifi_connection、performance、country_code）"],
        ["testcases/miniAP/", "miniAP 域用例"],
        ["testcases/smoke/", "冒烟验证，强调反馈速度，部分用例无需连接设备"],
        ["testcases/conftest.py", "全局 Fixture：device_serial、adb、wifi、wifi_yaml 等"],
        ["testcases/android/conftest.py", "Android 域 Fixture：android_wifi_lifecycle 自动前置/后置"],
    ])

    add_title(doc, "4.6 文档层 (docs/)", 2)
    add_bullets(doc, [
        "docs/testcase-docs/：测试用例说明文档的统一入口，与 testcases/ 分离。",
        "文档负责「可读性与追溯」，代码负责「可执行性」。",
        "用例 Python 文件头部注释可引用对应 md 源文档路径。",
    ])

    add_title(doc, "4.7 产出层 (reports/)", 2)
    add_para(doc, "输出路径：reports/YYYY-MM-DD/HH-MM-SS/，每次执行独立时间目录，避免同日覆盖。")
    add_bullets(doc, [
        "html/report.html — pytest-html 自包含 HTML 报告",
        "log/ — pytest.log 与 automation.log",
        "allure/raw/ — Allure 原始结果（需 --allure 或安装 allure-pytest 时自动写入）",
    ])

    # 5. 快速上手
    add_title(doc, "5. 快速上手指南", 1)
    add_title(doc, "5.1 环境准备", 2)
    add_code(
        doc,
        "# 1. 安装依赖\n"
        "pip install -r requirements.txt\n"
        "\n"
        "# 2. 配置环境变量（项目根目录 .env）\n"
        "ANDROID_SERIAL=your_device_serial\n"
        "WIFI_PASS_XXX=your_wifi_password\n"
        "PYAT_AI_API_KEY=optional\n"
        "\n"
        "# 3. 可选：config/test_config.json 填写 android.serial\n"
        "#    （仅当未设置 ANDROID_SERIAL 环境变量时生效）",
    )

    add_title(doc, "5.2 常用执行命令", 2)
    add_table(doc, ["命令", "说明"], [
        ["python run.py", "全量执行 testcases 目录"],
        ["python run.py --scene smoke", "冒烟场景（testcases/smoke + smoke marker）"],
        ["python run.py --scene hyperwifi", "Android 业务域（testcases/android）"],
        ["python run.py --scene stress", "Android 域 + stress marker（压力/长时）"],
        ["python run.py --scene weaknet", "Android 域 + weaknet marker（弱网）"],
        ["python run.py --scene miniap", "miniAP 业务域"],
        ["python run.py --scene ai", "全局 + ai marker"],
        ["python run.py --allure", "同时写入 Allure 原始结果"],
        ["python run.py --list", "仅收集用例，不执行（--collect-only）"],
        ["python run.py --scene smoke -q", "pytest 参数透传（-q 等追加到 pytest 命令末尾）"],
        ["pytest testcases -m \"wifi and not stress\"", "直连 pytest，灵活过滤"],
    ])

    # 6. 执行链路
    add_title(doc, "6. 执行链路（从命令到报告）", 1)
    steps = [
        "开发者通过 python run.py（或直连 pytest）发起测试。",
        "run.py 根据 --scene 确定收集目录与 marker 过滤策略。",
        "init_report_run_dir() 初始化 reports/<日期>/<时间>/ 目录。",
        "pytest 加载 pytest.ini 配置，收集 test_*.py 中的 test_* 函数。",
        "conftest.py 注入 Fixture（设备、ADB、WiFi 控制等），Android 域自动执行 WiFi 生命周期。",
        "用例执行，日志与步骤输出写入报告目录。",
        "pytest-html 生成自包含 HTML 报告；可选 Allure 原始数据。",
        "失败信息、日志、步骤证据统一落盘，形成可复盘资产。",
    ]
    for i, s in enumerate(steps, 1):
        add_para(doc, f"{i}. {s}")

    # 7. Marker 体系
    add_title(doc, "7. 标记体系 (Marker)", 1)
    add_para(doc, "pytest.ini 中 strict-markers 已启用，所有 marker 必须预先注册。主要标记如下：")
    add_table(doc, ["标记", "含义", "典型用途"], [
        ["wifi", "WiFi 产品/连接测试", "传统 WiFi 功能回归"],
        ["hyperwifi", "HyperWiFi 产品域", "HyperWiFi 专项测试"],
        ["wlan", "WLAN 业务顶层域", "新业务用例默认标记"],
        ["wlan_connection", "WLAN 连接子模块", "连接类用例细分"],
        ["wlan_country", "WLAN 国家码子模块", "国家码相关用例"],
        ["p0 ~ p3", "优先级", "需求文档优先级对齐"],
        ["regression", "回归测试", "全量/周期回归"],
        ["stress", "压力/长时测试", "长时间稳定性验证"],
        ["weaknet", "弱网仿真", "弱网场景专项"],
        ["smoke", "冒烟/快速 sanity", "CI 快速反馈"],
        ["miniap", "miniAP 产品域", "miniAP 设备测试"],
        ["ai", "调用 AI 模块", "AI 辅助用例生成/分析"],
        ["case_id", "用例 ID 追溯", "pytest.mark.case_id(\"WIFI-001\")"],
        ["no_wifi_lifecycle", "跳过 WiFi 前置/后置", "特殊场景保留 WiFi 状态"],
    ])
    add_para(doc, "Marker 价值：", bold=True)
    add_bullets(doc, [
        "业务维度切片：按产品域（wifi / hyperwifi / miniap）回归。",
        "风险维度切片：按压力/弱网等高风险场景拉通。",
        "执行策略切片：CI 中按时长和优先级拆分任务（快慢分离）。",
    ])

    # 8. Fixture 体系
    add_title(doc, "8. Fixture 体系", 1)
    add_title(doc, "8.1 全局 Fixture (testcases/conftest.py)", 2)
    add_table(doc, ["Fixture", "Scope", "说明"], [
        ["device_serial", "session", "从 DeviceManager 获取设备序列号，无设备则 skip"],
        ["adb", "session", "AdbHelper 实例，等待设备就绪"],
        ["wifi", "function", "WiFiControl 实例，每用例独立"],
        ["wifi_yaml", "session", "加载 testdata/android/wifi_ssid.yaml"],
        ["weaknet_yaml", "session", "加载 testdata/android/weaknet_scene.yaml"],
        ["run_shell_cmd", "function", "执行 adb shell 命令并返回文本"],
        ["ai_enabled", "function", "检查 PYAT_AI_API_KEY 是否配置"],
    ])

    add_title(doc, "8.2 Android 域生命周期 (android_wifi_lifecycle)", 2)
    add_para(
        doc,
        "testcases/android/conftest.py 中 autouse=True 的 Fixture，"
        "对使用 adb/wifi/run_shell_cmd 且带 WiFi 域 marker 的用例自动生效：",
    )
    add_bullets(doc, [
        "前置：检查并打开 WiFi 开关（确保可操作状态）。",
        "后置：断开所有 WiFi → 清除已保存网络 → 还原初始 WiFi 开关状态。",
        "纯 YAML/AI 数据用例（不依赖设备 Fixture）不启用。",
        "使用 @pytest.mark.no_wifi_lifecycle 可跳过生命周期。",
        "dual_wifi_eu / dual_wifi_us 等 marker 可控制双 WiFi 预置策略。",
    ])

    # 9. 核心能力详解
    add_title(doc, "9. 核心能力详解", 1)

    add_title(doc, "9.1 ADB 封装 (core/utils/adb_helper.py)", 2)
    add_para(
        doc,
        "AdbHelper 封装 adb 命令行调用，支持指定设备序列号、超时控制、"
        "shell 命令执行与设备等待。所有 Android 侧操作均通过此层与设备通信。",
    )

    add_title(doc, "9.2 WiFi 控制 (core/libs/android/wifi_control.py)", 2)
    add_para(doc, "WiFiControl 是 Android WiFi 自动化核心类，内部组合三个子模块：")
    add_bullets(doc, [
        "WiFiState：WiFi 开关状态查询与切换（cmd wifi / settings）。",
        "WiFiConnector：SSID 连接、等待关联、双 WiFi 连接等。",
        "WiFiCleanup：断开连接、清除已保存网络、等待断开确认。",
    ])
    add_para(doc, "主要能力：is_wifi_enabled、set_wifi_enabled、connect/disconnect、list_networks、clear_saved_networks 等。")

    add_title(doc, "9.3 数据加载 (core/testdata_loader/android/)", 2)
    add_para(
        doc,
        "wifi_data.py 提供 YAML 数据的结构化查询：get_wifi_list_entry（按 ssid/type/expect 筛选）、"
        "get_dual_wifi_entry（双 WiFi 配置）、get_expected_connected_ssids 等。"
        "用例通过 wifi_yaml Fixture 获取原始 dict，或调用 dataloader 辅助函数精确查询。",
    )

    add_title(doc, "9.4 报告路径管理 (core/utils/report_utils.py)", 2)
    add_bullets(doc, [
        "init_report_run_dir()：进程内固定本次运行目录，支持 PYTEST_REPORT_RUN_DIR 环境变量覆盖。",
        "REPORT_DATE / REPORT_TIME：可固定日期/时间子目录，便于调试或 CI 对齐。",
        "html_report_path() / allure_results_dir() / log_report_dir()：各子目录路径。",
        "conftest.py 的 pytest_configure 钩子自动将 HTML/Allure/日志写入当次目录。",
    ])

    add_title(doc, "9.5 日志 (core/utils/log_handler.py)", 2)
    add_para(
        doc,
        "统一日志格式：时间 | 级别 | 模块:文件:行号 | 消息。"
        "pytest 运行时不写 StreamHandler（避免与 capture 编码冲突），"
        "日志文件 UTF-8 写入 reports/<日期>/<时间>/log/ 目录。",
    )

    # 10. 用例编写规范
    add_title(doc, "10. 用例编写规范", 1)
    add_title(doc, "10.1 文件与函数命名", 2)
    add_bullets(doc, [
        "文件：test_*.py（pytest.ini 中 python_files 规则）。",
        "类：Test*（可选，当前项目以函数式用例为主）。",
        "函数：test_*，建议与 case_id 或需求 ID 对应，如 test_Perf_Conn_AssocTime_002。",
        "单文件用例数量建议 ≤100，可按模块继续拆分子目录。",
    ])

    add_title(doc, "10.2 典型用例结构", 2)
    add_code(
        doc,
        '@pytest.mark.p2\n'
        '@pytest.mark.wifi\n'
        '@pytest.mark.wlan_connection\n'
        '@pytest.mark.case_id("WiFi_Conn_Toggle_001")\n'
        'def test_WiFi_Conn_Toggle_001(run_shell_cmd, wifi: WiFiControl):\n'
        '    """\n'
        '    用例摘要：WiFi_Conn_Toggle_001\n'
        '    初始条件：设备 ADB 在线\n'
        '    测试目的：验证 WiFi 开关循环后 status 可正常查询\n'
        '    """\n'
        '    # Step 1: 打开 WiFi\n'
        '    assert wifi.set_wifi_enabled(True)\n'
        '    # Step 2: 关闭并查询\n'
        '    wifi.set_wifi_enabled(False)\n'
        '    out = run_shell_cmd("cmd wifi status")\n'
        '    assert out, "status 应有输出"',
    )

    add_title(doc, "10.3 性能用例模式", 2)
    add_para(
        doc,
        "性能类用例（testcases/android/performance/）使用 perf_case + run_*_case 辅助函数，"
        "统一封装拓扑描述、步骤日志与指标断言。源文档位于 docs/testcase-docs/android/performance/。",
    )

    add_title(doc, "10.4 实网 WiFi 注意事项", 2)
    add_bullets(doc, [
        "配置 testdata/android/wifi_ssid.yaml 与各 WIFI_PASS_* 环境变量（或 yaml 内 pwd 字段）。",
        "AP 不可达时对应用例会 skip 或 xfail，避免误报失败。",
        "Android 域 lifecycle 会在用例前后清理 WiFi 状态，注意与其他用例的隔离。",
    ])

    # 11. 依赖清单
    add_title(doc, "11. 主要依赖", 1)
    add_table(doc, ["包名", "用途"], [
        ["pytest (>=7,<8)", "测试执行引擎"],
        ["pytest-html", "HTML 报告生成"],
        ["pytest-xdist", "并行执行（可选）"],
        ["allure-pytest", "Allure 报告原始数据"],
        ["PyYAML", "测试数据 YAML 解析"],
        ["python-dotenv", ".env 环境变量加载"],
        ["requests / httpx", "HTTP 客户端（API/Web 测试、AI 调用）"],
        ["pandas / openpyxl", "数据导出与分析"],
        ["pillow", "截图/image 处理"],
    ])

    # 12. 现阶段收益
    add_title(doc, "12. 现阶段收益总结", 1)
    add_bullets(doc, [
        "目录职责清晰：代码、数据、文档、报告边界明确，新人可按层理解。",
        "场景执行统一：run.py 降低使用门槛，一条命令切换业务域与风险场景。",
        "结果可追溯：报告按日期/时间沉淀，HTML + 日志 + 可选 Allure 便于复盘。",
        "扩展成本可控：新增业务域可复用既有骨架（conftest + marker + dataloader 模式）。",
        "WiFi 生命周期自动化：Android 域前置/后置减少用例间状态污染。",
    ])

    # 13. 改进方向
    add_title(doc, "13. 当前可见改进点", 1)
    add_bullets(doc, [
        "miniAP 与 common 域能力可继续补齐，避免「目录已建但能力稀疏」。",
        "关键用例可补充 case_id 约束与文档映射，增强需求追踪闭环。",
        "失败诊断可继续标准化日志字段，提升跨人协作排障效率。",
        "回归任务可按 marker 做更细粒度的 CI 分层（快慢分离、风险分层）。",
        "pytest-xdist 并行策略可与设备池管理结合，提升多设备回归效率。",
    ])

    # 14. 团队协作建议
    add_title(doc, "14. 团队协作建议", 1)
    add_bullets(doc, [
        "新需求流程：先补文档 (docs/testcase-docs/) → 再落代码 (testcases/) → 最后更新数据 (testdata/)。",
        "新增工具能力优先落在 core/utils 或 core/libs，避免直接堆在测试脚本中。",
        "评审时同时看三件事：可执行性、可读性、可追溯性（case_id + 文档链接）。",
        "敏感信息（密码、序列号）只放 .env 或 CI 密钥，不提交到版本库。",
    ])

    # 15. 结论
    add_title(doc, "15. 汇报结论", 1)
    add_bullets(doc, [
        "我们已经形成「可运行、可维护、可扩展、可追溯」的自动化框架基线。",
        "当前框架对 Android WiFi / HyperWiFi 场景支撑成熟，对 miniAP 保持扩展位。",
        "下一阶段重点不是「从 0 到 1」，而是「从可用到高复用、高稳定、可度量」。",
    ])

    # Appendix
    add_title(doc, "附录 A：目录树速查", 1)
    add_code(
        doc,
        "PyAutoTest_IoT/\n"
        "├── run.py / pytest.ini / requirements.txt / README.md\n"
        "├── config/          test_config.json, env_config.json\n"
        "├── core/\n"
        "│   ├── utils/       adb_helper, config_loader, report_utils, log_handler ...\n"
        "│   ├── libs/android/ wifi_control, wifi_ui, wifi_cases ...\n"
        "│   ├── testdata_loader/ android/wifi_data, country_code_data\n"
        "│   └── ai/          AI 辅助模块（可选）\n"
        "├── testdata/android/    wifi_ssid.yaml, weaknet_scene.yaml, country_code_tables.yaml\n"
        "├── testcases/\n"
        "│   ├── conftest.py\n"
        "│   ├── android/     wifi_connection/, performance/, country_code/\n"
        "│   ├── miniAP/\n"
        "│   └── smoke/\n"
        "├── docs/            testcase-docs/, 本文档\n"
        "├── reports/         YYYY-MM-DD/HH-MM-SS/{html,log,allure}\n"
        "└── temp/            临时文件、截图",
    )

    add_title(doc, "附录 B：环境变量参考", 1)
    add_table(doc, ["变量名", "说明"], [
        ["ANDROID_SERIAL", "Android 设备序列号（优先于 test_config.json）"],
        ["ADB_TIMEOUT_SEC", "ADB 命令超时秒数，默认 30"],
        ["WIFI_PASS_*", "各 SSID 对应密码（与 wifi_ssid.yaml 配合）"],
        ["PYAT_AI_API_KEY", "AI 模块 API 密钥"],
        ["REPORT_DATE", "固定报告日期目录（如 2026-06-08）"],
        ["REPORT_TIME", "固定报告时间子目录（如 14-30-00）"],
        ["PYTEST_REPORT_RUN_DIR", "指定完整报告运行目录"],
        ["LOG_LEVEL", "日志级别，默认 INFO"],
    ])

    return doc


def main():
    doc = build_document()
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUTPUT))
    print(f"Generated: {OUTPUT}")


if __name__ == "__main__":
    main()
